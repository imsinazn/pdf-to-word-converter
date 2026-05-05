"""
PDF to Word Converter
A professional tool to convert PDF files to Word documents with images support.
Author: Sina
GitHub: https://github.com/imsinazn/pdf-to-word-converter
"""

import fitz  # PyMuPDF
from docx import Document
from docx.shared import Inches
import os

def convert_pdf_to_word(pdf_file, docx_file=None, start=0, end=None, keep_images=True):
    """
    Convert PDF to Word document with images support
    """
    if docx_file is None:
        docx_file = pdf_file.replace('.pdf', '.docx')
    
    try:
        # Check if input file exists
        if not os.path.exists(pdf_file):
            raise FileNotFoundError(f"File '{pdf_file}' not found!")
        
        print(f"\n🔄 Converting '{pdf_file}' to '{docx_file}'...")
        
        # Open PDF with PyMuPDF
        doc = fitz.open(pdf_file)
        total_pages = len(doc)
        
        # Set end page if not specified
        if end is None or end > total_pages:
            end = total_pages
        
        print(f"📖 Total pages: {total_pages}")
        
        # Ask user if they want specific pages
        specific_pages = input("\nDo you want specific pages only? (y/n): ").strip().lower()
        if specific_pages == 'y':
            try:
                start = int(input("Enter start page (0 for first): ") or 0)
                end_input = input("Enter end page (press Enter for last): ")
                end = int(end_input) if end_input else total_pages
                print(f"📖 Converting pages {start} to {end-1}")
            except ValueError:
                print("Invalid input. Converting all pages.")
                start, end = 0, total_pages
        else:
            start, end = 0, total_pages
            print(f"📖 Converting all {total_pages} pages")
        
        # Ask about images
        if keep_images:
            keep_images_input = input("\nExtract images from PDF? (y/n): ").strip().lower()
            keep_images = keep_images_input != 'n'
        
        # Create Word document
        word_doc = Document()
        images_extracted = 0
        pages_without_text = 0
        
        for page_num in range(start, end):
            print(f"📄 Processing page {page_num + 1}/{total_pages}...")
            
            # Get the page
            page = doc[page_num]
            
            # Extract and add text
            text = page.get_text()
            if text.strip():
                word_doc.add_paragraph(text)
            else:
                pages_without_text += 1
            
            # Extract and add images if requested
            if keep_images:
                image_list = page.get_images()
                
                for img_index, img in enumerate(image_list):
                    try:
                        # Get image data
                        xref = img[0]
                        base_image = doc.extract_image(xref)
                        image_bytes = base_image["image"]
                        image_ext = base_image["ext"]
                        
                        # Save image temporarily
                        temp_image_path = f"temp_img_{page_num}_{img_index}.{image_ext}"
                        with open(temp_image_path, "wb") as f:
                            f.write(image_bytes)
                        
                        # Add image to Word document
                        word_doc.add_picture(temp_image_path, width=Inches(5))
                        word_doc.add_paragraph("\n")
                        images_extracted += 1
                        
                        # Delete temp image
                        os.remove(temp_image_path)
                        
                    except Exception as e:
                        print(f"⚠️ Could not extract image on page {page_num}: {e}")
            
            # Add page break (except last page)
            if page_num < end - 1:
                word_doc.add_page_break()
        
        doc.close()
        
        # Save Word document
        word_doc.save(docx_file)
        
        file_size = os.path.getsize(docx_file)
        
        # Final report
        print("\n" + "="*50)
        print("✅ CONVERSION COMPLETE!")
        print("="*50)
        print(f"📁 Input file: {pdf_file}")
        print(f"📄 Output file: {docx_file}")
        print(f"📊 Statistics:")
        print(f"   - Pages processed: {end - start}")
        print(f"   - Images extracted: {images_extracted}")
        print(f"   - Pages without text: {pages_without_text}")
        print(f"   - File size: {file_size / 1024:.2f} KB")
        print("="*50)
        return True
        
    except FileNotFoundError as e:
        print(f"\n❌ Error: {e}")
        print("💡 Tip: Make sure the file name is correct and includes .pdf extension")
        return False
    except ImportError as e:
        print(f"\n❌ Missing library: {e}")
        print("🔧 Solution: Run: pip install PyMuPDF python-docx")
        return False
    except Exception as e:
        print(f"\n❌ Unexpected error: {e}")
        return False

def main():
    print("\n" + "="*50)
    print("📚 PDF TO WORD CONVERTER")
    print("="*50)
    
    while True:
        # Get filename from user
        pdf_file = input("\n📁 Enter PDF filename (with .pdf extension): ").strip()
        
        # Add .pdf if user forgot
        if pdf_file and not pdf_file.lower().endswith('.pdf'):
            pdf_file += '.pdf'
            print(f"ℹ️ Added .pdf extension: '{pdf_file}'")
        
        if not pdf_file:
            print("❌ No filename entered. Please try again.")
            continue
        
        # Try to convert
        success = convert_pdf_to_word(pdf_file)
        
        if success:
            # Ask if user wants to convert another file
            another = input("\n🔄 Convert another PDF? (y/n): ").strip().lower()
            if another != 'y':
                print("\n👋 Goodbye! Thanks for using PDF to Word Converter!")
                break
        else:
            retry = input("\n🔁 Try again with another filename? (y/n): ").strip().lower()
            if retry != 'y':
                print("\n👋 Goodbye!")
                break

if __name__ == "__main__":
    main()